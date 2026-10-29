# Tests de la hoja "Detalle por Guía" (3.ª hoja) del Excel del histórico
# (`POST /siscore/historico/exportar-excel`).
#
# La hoja baja un nivel más que "Detalle por Pedido": el flete asignado a cada
# carro se reparte entre sus guías por piezas (mismo _repartir_flete) → la suma
# de las guías de un carro es EXACTAMENTE su costo asignado, y la del total el
# total_solicitado del vehículo. Docs viejos sin registros_detalle usan las
# guías del campo codigo_pedido (reparto equitativo); carro sin guías conserva
# su costo en una fila "-".
import asyncio
import io
import unittest
from datetime import datetime, timezone
from unittest.mock import patch

from openpyxl import load_workbook

from rutas import siscore_consultas as mod
from rutas.siscore_consultas import ExportarHistoricoExcelRequest, exportar_historico_excel


class FakeCursor(list):
    """Cursor encadenable: .sort(...) devuelve las filas tal cual (orden irrelevante)."""

    def sort(self, *args, **kwargs):
        return self


class FakeCollection:
    """MongoCollection en memoria: devuelve todo (el test controla los docs)."""

    def __init__(self, docs=None):
        self.docs = list(docs or [])

    def find(self, query=None, projection=None):
        return FakeCursor(dict(d) for d in self.docs)

    def find_one(self, query=None):
        return self.docs[0] if self.docs else None


def _doc_fusion():
    """Fusión de 2 carros; cada original trae sus registros_detalle por guía."""
    return {
        "consecutivo": "FUNZA-20261028-1A",
        "planilla": "824986",
        "total_solicitado": 1000000,
        "estado": "APROBADO",
        "cliente_origen": "FRESENIUS MEDICAL CARE",
        "fecha_creacion": datetime(2026, 10, 28, tzinfo=timezone.utc),
        "fusion_info": {
            "es_fusionada": True,
            "datos_originales": [
                {
                    "consecutivo": "FUNZA-20261028-1A",
                    "planilla": "824986",
                    "pedido_vulcano": "120795",
                    "cliente_origen": "FRESENIUS MEDICAL CARE",
                    "piezas": 10,
                    "peso_real": 100,
                    "registros_detalle": [
                        {"Guia": "9001", "Cliente Origen": "FRESENIUS MEDICAL CARE",
                         "Nombre": "JUAN PEREZ", "Municipio Destino": "BOGOTA",
                         "Piezas": 6, "Peso Real": 60},
                        {"Guia": "9002", "Cliente Origen": "FRESENIUS MEDICAL CARE",
                         "Nombre": "MARIA LOPEZ", "Municipio Destino": "SOACHA",
                         "Piezas": 4, "Peso Real": 40},
                    ],
                },
                {
                    "consecutivo": "FUNZA-20261028-1B",
                    "planilla": "824987",
                    "pedido_vulcano": "120796",
                    "cliente_origen": "FRESENIUS KABI",
                    "piezas": 5,
                    "peso_real": 50,
                    "registros_detalle": [
                        {"Guia": "9003", "Cliente Origen": "FRESENIUS KABI",
                         "Nombre": "CLINICA NORTE", "Municipio Destino": "MADRID",
                         "Piezas": 5, "Peso Real": 50},
                    ],
                },
            ],
        },
    }


def _doc_normal_viejo():
    """Planilla no fusionada SIN registros_detalle (doc viejo): fallback a
    codigo_pedido con reparto equitativo."""
    return {
        "consecutivo": "CALI-20261028-2",
        "planilla": "824990",
        "pedido_vulcano": "120800",
        "codigo_pedido": "9004, 9005",
        "total_solicitado": 300000,
        "estado": "APROBADO",
        "cliente_origen": "FRESENIUS KABI",
        "fecha_creacion": datetime(2026, 10, 28, tzinfo=timezone.utc),
    }


def _generar_excel(docs):
    historico = FakeCollection(docs)
    usuarios = FakeCollection([])
    with patch.object(mod, "coleccion_historico", historico), \
         patch.object(mod, "coleccion_baseusuarios", usuarios):
        resp = asyncio.run(exportar_historico_excel(ExportarHistoricoExcelRequest(
            fecha_inicio="2026-10-01", fecha_fin="2026-10-31",
            perfil="ADMIN", centro_distribucion=None,
        )))
    return load_workbook(io.BytesIO(resp.body))


def _filas_hoja(wb, nombre):
    ws = wb[nombre]
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    # La última fila puede ser TOTALES
    if filas and filas[-1][0] == "TOTALES":
        return filas[:-1], filas[-1]
    return filas, None


class TestHojaDetallePorGuia(unittest.TestCase):

    def test_hoja_existe_y_suma_exacta(self):
        wb = _generar_excel([_doc_fusion(), _doc_normal_viejo()])
        self.assertEqual(wb.sheetnames, ["Historico Pedidos", "Detalle por Pedido", "Detalle por Guía"])

        filas, totales = _filas_hoja(wb, "Detalle por Guía")
        # 3 guías de la fusión + 2 del doc viejo (codigo_pedido)
        self.assertEqual(len(filas), 5)

        # Carro A: 1M * 10/15 = 666.666; guía 9001 = 666.666 * 6/10 = 399.999
        f9001 = next(f for f in filas if f[4] == "9001")
        self.assertEqual(f9001[1], "FUNZA-20261028-1A")
        self.assertEqual(f9001[3], "120795")
        self.assertEqual(f9001[10], 399999)

        # Guía 9002 absorbe el residuo del carro A: 666.666 - 399.999 = 266.667
        f9002 = next(f for f in filas if f[4] == "9002")
        self.assertEqual(f9002[10], 266667)

        # Carro B absorbe el residuo de la fusión: 1M - 666.666 = 333.334
        f9003 = next(f for f in filas if f[4] == "9003")
        self.assertEqual(f9003[1], "FUNZA-20261028-1B")
        self.assertEqual(f9003[10], 333334)

        # Doc viejo: 300000 equitativo entre 9004 y 9005 (150000 c/u)
        f9004 = next(f for f in filas if f[4] == "9004")
        f9005 = next(f for f in filas if f[4] == "9005")
        self.assertEqual(f9004[10], 150000)
        self.assertEqual(f9005[10], 150000)

        # La suma de TODAS las guías es exactamente el total de cada vehículo
        total_fusion = sum(f[10] for f in filas if f[0] == "FUNZA-20261028-1A")
        self.assertEqual(total_fusion, 1000000)
        total_viejo = sum(f[10] for f in filas if f[0] == "CALI-20261028-2")
        self.assertEqual(total_viejo, 300000)

        # Fila TOTALES global = 1.300.000
        self.assertIsNotNone(totales)
        self.assertEqual(totales[10], 1300000)

    def test_carro_sin_guias_conserva_costo(self):
        doc = _doc_fusion()
        # El carro B pierde su detalle y su codigo_pedido → fila "-" con su flete
        carro_b = doc["fusion_info"]["datos_originales"][1]
        carro_b["registros_detalle"] = []
        carro_b.pop("codigo_pedido", None)

        wb = _generar_excel([doc])
        filas, totales = _filas_hoja(wb, "Detalle por Guía")
        fila_guion = next(f for f in filas if f[4] == "-")
        self.assertEqual(fila_guion[1], "FUNZA-20261028-1B")
        self.assertEqual(fila_guion[10], 333334)
        self.assertEqual(totales[10], 1000000)

    def test_coherencia_con_hoja_detalle_por_pedido(self):
        """El costo asignado a cada carro en la hoja 2 debe ser la suma exacta
        de sus guías en la hoja 3."""
        wb = _generar_excel([_doc_fusion(), _doc_normal_viejo()])
        filas_pedido, _ = _filas_hoja(wb, "Detalle por Pedido")
        filas_guia, _ = _filas_hoja(wb, "Detalle por Guía")

        for fp in filas_pedido:
            cons_orig = fp[1]
            suma_guias = sum(f[10] for f in filas_guia if f[1] == cons_orig)
            self.assertEqual(suma_guias, fp[8],
                             f"Carro {cons_orig}: guías suman {suma_guias} vs pedido {fp[8]}")


if __name__ == "__main__":
    unittest.main()
