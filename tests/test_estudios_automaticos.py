"""Tests de los estudios de seguridad automáticos (TusDatos) al pasar
un vehículo a completado_revision.

Cubre:
- sujetos_estudio: deduplicación de personas (roles combinados) + vehículo.
- disparar_estudios: escritura del array, re-revisión (huella), sin
  credenciales (configuracion_faltante), placa inválida, error del handler.
- Hook en actualizar-estado: dispara SOLO al llegar a completado_revision.
- GET /vehiculos/estudios-seguridad/{placa}.
- Blindajes: CLAVES_PROTEGIDAS / CAMPOS_VOLATILES_FIRMA / pop en listados.
"""
import sys
import types
import unittest
from datetime import datetime, timedelta
from unittest.mock import MagicMock, patch

# Atlas es inalcanzable intermitentemente desde la red local (SRV/DNS) y
# estos tests NO necesitan Mongo real (toda colección se parchea por test).
# Si el módulo de conexión no está cargado aún, se substituye por un stub;
# si ya está cargado (suite completa con red viva), no se toca.
if "bd.bd_cliente" not in sys.modules:
    _stub = types.ModuleType("bd.bd_cliente")

    class _BDFake(dict):
        def __getitem__(self, clave):
            return self.setdefault(clave, MagicMock())

    _stub.bd_cliente = _BDFake()
    sys.modules["bd.bd_cliente"] = _stub

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from rutas import vehiculos
from Funciones import estudios_automaticos


def cliente_de_prueba() -> TestClient:
    """Monta el router en una mini-app (patrón test_obligatoriedad_documentos)."""
    app = FastAPI()
    app.include_router(vehiculos.ruta_vehiculos)
    return TestClient(app, raise_server_exceptions=False)


class FakeColeccion:
    """Falsa colección vehiculos con soporte de $set punteado y
    array_filters sobre estudiosSeguridadAuto.$[e]."""

    def __init__(self, documentos=None):
        self.documents = list(documentos or [])
        self.updates = []

    def find_one(self, query, *args, **kwargs):
        for d in self.documents:
            if d.get("placa") == query.get("placa"):
                return d
        return None

    def find(self, filtro=None, *args, **kwargs):
        def _get_punteado(doc, clave):
            actual = doc
            for parte in clave.split("."):
                if not isinstance(actual, dict):
                    return None
                actual = actual.get(parte)
            return actual

        def match(d):
            for k, v in (filtro or {}).items():
                valor = _get_punteado(d, k)
                if isinstance(v, dict):
                    if "$in" in v and valor not in v["$in"]:
                        return False
                    if "$lt" in v and not (valor is not None and valor < v["$lt"]):
                        return False
                    if "$exists" in v and (valor is not None) != bool(v["$exists"]):
                        return False
                    if "$elemMatch" in v:
                        arr = valor if isinstance(valor, list) else []
                        cond = v["$elemMatch"]
                        if not any(
                            isinstance(e, dict) and
                            all(_get_punteado(e, k2) == v2 for k2, v2 in cond.items())
                            for e in arr
                        ):
                            return False
                elif valor != v:
                    return False
            return True
        return [d for d in self.documents if match(d)]

    @staticmethod
    def _set_punteado(doc: dict, clave: str, valor):
        partes = clave.split(".")
        actual = doc
        for p in partes[:-1]:
            actual = actual.setdefault(p, {})
        actual[partes[-1]] = valor

    def update_one(self, filtro, cambio, array_filters=None):
        self.updates.append((filtro, cambio, array_filters))
        for d in self.documents:
            if d.get("placa") != filtro.get("placa"):
                continue
            if "$set" in cambio:
                for k, v in cambio["$set"].items():
                    if k.startswith("estudiosSeguridadAuto.$[e].") and array_filters:
                        campo = k.split(".$[e].", 1)[1]
                        cond = array_filters[0].get("e.id")
                        for e in d.get("estudiosSeguridadAuto") or []:
                            if e.get("id") == cond:
                                e[campo] = v
                    elif k == "estudiosSeguridadAuto":
                        d["estudiosSeguridadAuto"] = v
                    else:
                        self._set_punteado(d, k, v)
            if "$push" in cambio:
                for k, v in cambio["$push"].items():
                    if isinstance(v, dict) and "$each" in v:
                        actual = d.setdefault(k, [])
                        if v.get("$position") == 0:
                            actual[0:0] = v["$each"]
                        else:
                            actual.extend(v["$each"])
                        tope = v.get("$slice")
                        if isinstance(tope, int) and tope > 0:
                            del actual[tope:]
                    else:
                        d.setdefault(k, []).append(v)


TODOS_DOCS = {
    "tarjetaPropiedad": "https://x/1", "tarjetaPropiedadReverso": "https://x/1r",
    "soat": "https://x/2", "revisionTecnomecanica": "https://x/3",
    "documentoIdentidadConductor": "https://x/6",
    "documentoIdentidadConductorReverso": "https://x/6r",
    "documentoIdentidadPropietario": "https://x/7",
    "documentoIdentidadPropietarioReverso": "https://x/7r",
    "documentoIdentidadTenedor": "https://x/8",
    "documentoIdentidadTenedorReverso": "https://x/8r",
    "licencia": "https://x/9", "licenciaReverso": "https://x/9r",
    "planillaEpsArl": "https://x/10", "condFoto": "https://x/11",
    "condCertificacionBancaria": "https://x/12",
    "tenedCertificacionBancaria": "https://x/14",
    "documentoAcreditacionTenedor": "https://x/15",
    "rutTenedor": "https://x/16", "fotos": ["https://x/f1"],
    "hojaVidaFisica": "https://x/17",
}


def vehiculo_completo(**extra):
    doc = {
        "_id": "id-prueba",
        "placa": "ABC123", "estadoIntegra": "registro_incompleto",
        "vehCapacidadCarga": "8000", **TODOS_DOCS,
    }
    doc.update(extra)
    return doc


async def _respuesta_ok(*args, **resultado_extra):
    """Respuesta feliz del wrapper (consulta_completa/consulta_vehiculo).
    Recibe el modelo pydantic como argumento posicional."""
    resultado = {
        "estado": "finalizado", "hallazgo": False, "hallazgos": "",
        "results": {"SIMIT": False, "PROCURADURIA": "Error"}, "id": "rep-123",
    }
    resultado.update(resultado_extra)
    return {"lanzamiento": {"jobid": "job-1", "nombre": "FULANO PEREZ"},
            "resultado": resultado}


class SujetosEstudioTests(unittest.TestCase):

    def test_dedup_combina_roles(self):
        sujetos = estudios_automaticos.sujetos_estudio({
            "condCedulaCiudadania": "1.020.304.050",
            "propDocumento": "1020304050",   # mismo que conductor (dedup)
            "tenedDocumento": "987.654.321",
            "placa": "ABC123",
        })
        personas = [s for s in sujetos if s["tipo"] == "persona"]
        vehiculo = [s for s in sujetos if s["tipo"] == "vehiculo"]
        self.assertEqual(len(personas), 2)
        self.assertEqual(personas[0]["roles"], ["conductor", "propietario"])
        self.assertEqual(personas[0]["cedula"], "1020304050")
        self.assertEqual(personas[1]["roles"], ["tenedor"])
        self.assertEqual(len(vehiculo), 1)
        self.assertEqual(vehiculo[0]["placa"], "ABC123")
        self.assertEqual(vehiculo[0]["cedula_propietario"], "1020304050")

    def test_los_tres_iguales_un_solo_estudio_de_persona(self):
        sujetos = estudios_automaticos.sujetos_estudio({
            "condCedulaCiudadania": "1020304050",
            "propDocumento": "1020304050", "tenedDocumento": "1020304050",
            "placa": "ABC123",
        })
        personas = [s for s in sujetos if s["tipo"] == "persona"]
        self.assertEqual(len(personas), 1)
        self.assertEqual(personas[0]["roles"], ["conductor", "propietario", "tenedor"])

    def test_propietario_nit_genera_sujeto_empresa(self):
        sujetos = estudios_automaticos.sujetos_estudio({
            "condCedulaCiudadania": "1020304050",
            "propDocumento": "901923029", "propTipoDocumento": "NIT",
            "tenedDocumento": "987654321",
            "placa": "ABC123",
        })
        tipos = [(s["tipo"], s.get("cedula") or s.get("nit") or s.get("placa")) for s in sujetos]
        self.assertIn(("empresa", "901923029"), tipos)
        self.assertNotIn(("persona", "901923029"), tipos)  # NO como cédula
        empresa = [s for s in sujetos if s["tipo"] == "empresa"][0]
        self.assertEqual(empresa["roles"], ["propietario"])

    def test_estudio_de_empresa_usa_el_flujo_nit(self):
        """El sujeto empresa consulta por launch/verify/nit + sondeo, y su
        PDF se archiva con el endpoint pdf-nit."""
        veh = vehiculo_completo(
            condCedulaCiudadania="1020304050",
            propDocumento="901923029", propTipoDocumento="NIT")
        fake = FakeColeccion([veh])

        async def lanzar_nit(nit):
            assert int(nit.nit) == 901923029
            return {"jobid": "job-nit"}

        async def esperar(jobid, maximo, intervalo):
            assert jobid == "job-nit"
            return {"estado": "finalizado", "hallazgo": False, "hallazgos": "",
                    "results": {"RUES": False}, "id": "rep-nit-1"}

        async def pdf_nit(_id):
            return b"%PDF-empresa"

        subidas = []
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", _respuesta_ok), \
             patch.object(estudios_automaticos, "_td_verificar_nit", lanzar_nit), \
             patch.object(estudios_automaticos, "_td_esperar", esperar), \
             patch.object(estudios_automaticos, "_td_reporte_nit_pdf", pdf_nit), \
             patch.object(estudios_automaticos, "_subir_blob_pdf",
                          lambda ruta, c: subidas.append(ruta)):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
        empresa = [e for e in fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"]
                   if e["tipo"] == "empresa"][0]
        self.assertEqual(empresa["estado"], "finalizado")
        self.assertEqual(empresa["reporte_id"], "rep-nit-1")
        self.assertTrue(any("empresa_901923029" in r for r in subidas))

    def test_documentos_propietario_empresa(self):
        """Con propietario NIT: la cédula del propietario NO se exige y el
        RUT de la empresa SÍ (espejo del conductor en el front)."""
        # Propietario empresa SIN cédula y SIN RUT → falta rutPropietario.
        veh = vehiculo_completo(propDocumento="901923029",
                                propTipoDocumento="NIT")
        faltan = vehiculos._documentos_faltantes(veh)
        self.assertNotIn("documentoIdentidadPropietario", faltan)
        self.assertNotIn("documentoIdentidadPropietarioReverso", faltan)
        self.assertIn("rutPropietario", faltan)
        # Con el RUT subido → completo.
        veh["rutPropietario"] = "Vehiculos/ABC123/rut.pdf"
        self.assertEqual(vehiculos._documentos_faltantes(veh), [])

    def test_cedula_vacia_no_genera_sujeto(self):
        sujetos = estudios_automaticos.sujetos_estudio({
            "condCedulaCiudadania": "", "placa": "ABC123"})
        self.assertEqual([s["tipo"] for s in sujetos], ["vehiculo"])

    def test_fecha_expedicion_del_conductor_viaja_al_sujeto(self):
        sujetos = estudios_automaticos.sujetos_estudio({
            "condCedulaCiudadania": "1020304050",
            "condFechaExpedicion": "2012-02-14",  # ISO, como la guarda la IA
            "propDocumento": "1020304050", "tenedDocumento": "987654321",
            "placa": "ABC123"})
        por_cedula = {s["cedula"]: s for s in sujetos if s["tipo"] == "persona"}
        self.assertEqual(por_cedula["1020304050"]["fecha_expedicion"], "14/02/2012")
        self.assertEqual(por_cedula["987654321"]["fecha_expedicion"], "")  # sin fecha

    def test_tdoc_del_vehiculo_se_deriva_del_rut(self):
        sujetos = estudios_automaticos.sujetos_estudio({
            "condCedulaCiudadania": "1020304050",
            "propDocumento": "901923029", "propTipoDocumento": "NIT",
            "placa": "ABC123"})
        vehiculo = [s for s in sujetos if s["tipo"] == "vehiculo"][0]
        self.assertEqual(vehiculo["tdoc_propietario"], "NIT")
        # Sin tipo de documento → CC (persona natural)
        sujetos2 = estudios_automaticos.sujetos_estudio({
            "condCedulaCiudadania": "1020304050", "propDocumento": "1020304050",
            "placa": "ABC123"})
        self.assertEqual(sujetos2[-1]["tdoc_propietario"], "CC")


class DispararEstudiosTests(unittest.TestCase):
    """disparar_estudios corre la ejecución completa (await directo):
    asyncio.run espera los updates del handler falso."""

    def _correr(self, fake, vehiculo):
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", _respuesta_ok), \
             patch.object(estudios_automaticos, "consulta_vehiculo", _respuesta_ok):
            asyncio.run(estudios_automaticos.disparar_estudios(vehiculo["placa"]))
        return fake.find_one({"placa": vehiculo["placa"]})

    def test_flujo_feliz_persiste_resultado(self):
        fake = FakeColeccion([vehiculo_completo(
            condCedulaCiudadania="1020304050",
            propDocumento="1020304050", tenedDocumento="987654321")])
        doc = self._correr(fake, fake.documents[0])
        estudios = doc["estudiosSeguridadAuto"]
        self.assertEqual(len(estudios), 3)  # persona combinada + tenedor + vehículo
        for e in estudios:
            self.assertEqual(e["estado"], "finalizado")
            self.assertEqual(e["proveedor"], "tusdatos")
            self.assertEqual(e["reporte_id"], "rep-123")
            self.assertEqual(e["categoria"], "")
            self.assertIn("fuentes", e)
            self.assertIn("finalizado_en", e)
        combinada = [e for e in estudios if e.get("roles")][0]
        self.assertEqual(combinada["roles"], ["conductor", "propietario"])

    def test_re_revision_sin_cambio_no_re_dispara(self):
        veh = vehiculo_completo(
            condCedulaCiudadania="1020304050",
            propDocumento="1020304050", tenedDocumento="987654321",
            estudiosSeguridadAuto=[{  # corrida previa con la misma huella
                "id": "aaa", "tipo": "persona", "cedula": "1020304050",
                "roles": ["conductor", "propietario"], "estado": "finalizado",
            }, {
                "id": "bbb", "tipo": "persona", "cedula": "987654321",
                "roles": ["tenedor"], "estado": "finalizado",
            }, {
                "id": "ccc", "tipo": "vehiculo", "placa": "ABC123",
                "estado": "finalizado",
            }])
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123", re_revision=True))
        self.assertEqual(fake.updates, [])  # conserva los estudios anteriores

    def test_re_revision_con_cedula_cambiada_re_dispara(self):
        veh = vehiculo_completo(
            condCedulaCiudadania="999999999",  # cambió respecto de la previa
            propDocumento="999999999", tenedDocumento="987654321",
            estudiosSeguridadAuto=[{
                "id": "aaa", "tipo": "persona", "cedula": "1020304050",
                "roles": ["conductor", "propietario"], "estado": "finalizado",
            }])
        fake = FakeColeccion([veh])
        doc = self._correr(fake, veh)
        cedulas = [e.get("cedula") for e in doc["estudiosSeguridadAuto"]
                   if e["tipo"] == "persona"]
        self.assertIn("999999999", cedulas)

    def test_placa_de_7_no_tumba_las_personas(self):
        veh = vehiculo_completo(
            placa="ABC1234",  # el registro acepta 4-7; el wrapper solo 6
            condCedulaCiudadania="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        doc = self._correr(fake, veh)
        por_tipo = {e["tipo"]: e for e in doc["estudiosSeguridadAuto"]}
        self.assertEqual(por_tipo["vehiculo"]["estado"], "error")
        self.assertEqual(por_tipo["vehiculo"]["error"], "placa_invalida")
        # Las personas terminaron bien aunque el vehículo fallara.
        personas_ok = all(e["estado"] == "finalizado"
                          for e in doc["estudiosSeguridadAuto"]
                          if e["tipo"] == "persona")
        self.assertTrue(personas_ok)

    def test_sin_credenciales_marca_configuracion_faltante(self):
        veh = vehiculo_completo(
            condCedulaCiudadania="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: False):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
        estudios = fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"]
        self.assertEqual(len(estudios), 3)
        for e in estudios:
            self.assertEqual(e["estado"], "error")
            self.assertEqual(e["error"], "configuracion_faltante")

    def test_handler_que_lanza_deja_el_estudio_en_error(self):
        async def explota(*a, **kw):
            raise HTTPException(status_code=502, detail="upstream caído")

        veh = vehiculo_completo(
            condCedulaCiudadania="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", explota), \
             patch.object(estudios_automaticos, "consulta_vehiculo", explota):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
        estudios = fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"]
        self.assertEqual(len(estudios), 3)
        for e in estudios:
            self.assertEqual(e["estado"], "error")
            self.assertIn("upstream caído", e["error"])

    def test_reintentar_fuentes_actualiza_el_estudio(self):
        """retry del proveedor → jobid → sondeo → persiste el mapa nuevo."""
        estudio = {
            "id": "e1", "tipo": "persona", "cedula": "1020304050",
            "roles": ["conductor"], "estado": "finalizado",
            "fuentes": {"SIMIT": "Error", "ONU": False}, "reporte_id": "rep-1",
        }
        veh = vehiculo_completo(estudiosSeguridadAuto=[estudio])
        fake = FakeColeccion([veh])

        async def retry(*args, **kwargs):
            return {"jobid": "job-retry"}

        async def esperar(jobid, maximo, intervalo):
            return {"estado": "finalizado", "hallazgo": False, "hallazgos": "",
                    "results": {"SIMIT": False, "ONU": False}, "id": "rep-1"}

        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_td_reintentar", retry), \
             patch.object(estudios_automaticos, "_td_esperar", esperar):
            cambios = asyncio.run(
                estudios_automaticos.reintentar_fuentes_estudio("ABC123", "e1"))
        self.assertEqual(cambios["fuentes"], {"SIMIT": False, "ONU": False})
        actualizado = fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"][0]
        self.assertEqual(actualizado["fuentes"]["SIMIT"], False)  # ya sin Error

    def test_reintentar_fuentes_de_VEHICULO_no_soportado(self):
        """Regresión (2026-10-07, QYO235): /api/retry/{id} responde {} para los
        ids de reportes de vehículo (probado en vivo) — antes eso caía al 502
        engañoso "estado=desconocido". Ahora: 422 claro SIN llamar al
        proveedor; la salida es «Volver a consultar»."""
        from fastapi import HTTPException
        estudio = {
            "id": "e1", "tipo": "vehiculo", "placa": "ABC123",
            "estado": "finalizado", "fuentes": {"Omisiones Vehículos": "Error"},
            "reporte_id": "rep-car",
        }
        veh = vehiculo_completo(estudiosSeguridadAuto=[estudio])
        fake = FakeColeccion([veh])

        async def retry(*a, **k):
            raise AssertionError("No debe llamarse al proveedor para vehículos")

        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_td_reintentar", retry):
            with self.assertRaises(HTTPException) as ctx:
                asyncio.run(estudios_automaticos.reintentar_fuentes_estudio("ABC123", "e1"))
        self.assertEqual(ctx.exception.status_code, 422)
        self.assertIn("VEHÍCULO", ctx.exception.detail)

    def test_reintentar_respuesta_vacia_del_proveedor_es_422_claro(self):
        """El proveedor puede rechazar el retry en silencio ({}): antes se
        transformaba en 502 "estado=desconocido" — ahora 422 accionable."""
        from fastapi import HTTPException
        estudio = {
            "id": "e1", "tipo": "persona", "cedula": "1020304050",
            "estado": "finalizado", "fuentes": {"SIMIT": "Error"}, "reporte_id": "rep-1",
        }
        veh = vehiculo_completo(estudiosSeguridadAuto=[estudio])
        fake = FakeColeccion([veh])

        async def retry(*a, **k):
            return {}

        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_td_reintentar", retry):
            with self.assertRaises(HTTPException) as ctx:
                asyncio.run(estudios_automaticos.reintentar_fuentes_estudio("ABC123", "e1"))
        self.assertEqual(ctx.exception.status_code, 422)
        self.assertIn("vacío", ctx.exception.detail)

    def test_reintentar_empresa_usa_typedoc_nit(self):
        estudio = {
            "id": "e1", "tipo": "empresa", "nit": "900123456",
            "estado": "finalizado", "fuentes": {"RUES": "Error"}, "reporte_id": "rep-nit",
        }
        veh = vehiculo_completo(estudiosSeguridadAuto=[estudio])
        fake = FakeColeccion([veh])
        llamadas = {}

        async def retry(id_reporte, typedoc="CC"):
            llamadas["typedoc"] = typedoc
            return {"jobid": "job-nit"}

        async def esperar(jobid, maximo, intervalo):
            return {"estado": "finalizado", "hallazgo": False, "hallazgos": "",
                    "results": {"RUES": False}, "id": "rep-nit"}

        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_td_reintentar", retry), \
             patch.object(estudios_automaticos, "_td_esperar", esperar):
            asyncio.run(estudios_automaticos.reintentar_fuentes_estudio("ABC123", "e1"))
        self.assertEqual(llamadas["typedoc"], "NIT")

    def test_pdf_del_reporte_se_archiva_en_el_bucket(self):
        """Al finalizar, el estudio persiste pdf_gcs con la ruta del blob
        (best-effort: si la subida falla, el estudio IGUAL finaliza)."""
        subidas = []

        async def pdf(_id):
            return b"%PDF-1.4 contenido de prueba"

        def subir(ruta, contenido):
            subidas.append((ruta, contenido))

        veh = vehiculo_completo(
            condCedulaCiudadaria="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", _respuesta_ok), \
             patch.object(estudios_automaticos, "consulta_vehiculo", _respuesta_ok), \
             patch.object(estudios_automaticos, "_td_reporte_pdf", pdf), \
             patch.object(estudios_automaticos, "_td_reporte_car_pdf", pdf), \
             patch.object(estudios_automaticos, "_subir_blob_pdf", subir):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
        estudios = fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"]
        self.assertEqual(len(subidas), 3)  # un PDF por estudio finalizado
        for e in estudios:
            self.assertEqual(e["estado"], "finalizado")
            self.assertIn("pdf_gcs", e)
            self.assertTrue(e["pdf_gcs"]["ruta"].startswith("Vehiculos/ABC123/"))
            # La placa va AL FINAL del nombre (2026-10-02)
            self.assertTrue(e["pdf_gcs"]["ruta"].endswith("_rep-123_abc123.pdf"))
        # La ruta es determinística por reporte_id (persona por cédula)
        self.assertIn("persona_1020304050_rep-123", subidas[0][0])

    def test_fallo_del_archivo_no_tumba_el_estudio(self):
        def subir(_ruta, _contenido):
            raise RuntimeError("GCS caído")

        veh = vehiculo_completo(
            condCedulaCiudadaria="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", _respuesta_ok), \
             patch.object(estudios_automaticos, "consulta_vehiculo", _respuesta_ok), \
             patch.object(estudios_automaticos, "_subir_blob_pdf", subir):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
        estudios = fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"]
        for e in estudios:
            self.assertEqual(e["estado"], "finalizado")  # archivado falló y siguió
            self.assertNotIn("pdf_gcs", e)

    def test_handlers_pdf_devuelven_bytes(self):
        """Regresión (2026-10-02, WOO453): los handlers de PDF que consume el
        módulo deben devolver los BYTES del reporte — antes se importaban los
        handlers de RUTA de /tusdatos (Response de FastAPI, sin .content) y
        tanto el archivado como el endpoint /pdf/ reventaban con
        'Response' object has no attribute 'content'."""
        from rutas import tusdatos

        class RespuestaHttpx:
            content = b"%PDF-real"

        rutas_llamadas = []

        async def pedido(_metodo, ruta, **_kw):
            rutas_llamadas.append(ruta)
            return RespuestaHttpx()

        import asyncio
        with patch.object(tusdatos, "_pedido", pedido):
            persona = asyncio.run(estudios_automaticos._td_reporte_pdf("rep-1"))
            nit = asyncio.run(estudios_automaticos._td_reporte_nit_pdf("rep-2"))
            car = asyncio.run(estudios_automaticos._td_reporte_car_pdf("rep-3"))
        self.assertEqual(persona, b"%PDF-real")
        self.assertEqual(nit, b"%PDF-real")
        self.assertEqual(car, b"%PDF-real")
        # Cada tipo de reporte va a SU endpoint del proveedor.
        self.assertIn("/api/v2/report_pdf/rep-1", rutas_llamadas)
        self.assertIn("/api/v2/report_nit_pdf/rep-2", rutas_llamadas)
        self.assertIn("/api/v2/report_car_pdf/rep-3", rutas_llamadas)

    def test_pdf_de_vehiculo_usa_el_endpoint_car(self):
        """Regresión (2026-10-07, QYO235): los estudios de VEHÍCULO
        (launch/car) se archivan y abren con /report_car_pdf — antes iban por
        /report_pdf (persona) y el proveedor respondía 410 «identificador
        inválido», disfrazado de «expiró». El archivado fallaba EN SILENCIO
        desde siempre (best-effort)."""
        import asyncio
        llamados = {"persona": 0, "car": 0}

        async def pdf_persona(_id):
            llamados["persona"] += 1
            return b"%PDF-persona"

        async def pdf_car(_id):
            llamados["car"] += 1
            return b"%PDF-car"

        subidas = []

        def subir(ruta, contenido):
            subidas.append((ruta, contenido))

        with patch.object(estudios_automaticos, "_td_reporte_pdf", pdf_persona), \
             patch.object(estudios_automaticos, "_td_reporte_car_pdf", pdf_car), \
             patch.object(estudios_automaticos, "_subir_blob_pdf", subir):
            resultado = asyncio.run(estudios_automaticos._archivar_pdf_reporte(
                "QYO235", {"tipo": "vehiculo", "placa": "QYO235"}, "rep-car-9"))
        self.assertIsNotNone(resultado)
        self.assertEqual(llamados, {"persona": 0, "car": 1})
        self.assertIn("vehiculo_QYO235_rep-car-9", subidas[0][0])
        self.assertEqual(subidas[0][1], b"%PDF-car")

    def test_vigencia_se_sella_al_disparar(self):
        """La corrida lleva estudiosVigencia {desde, vence=+VIGENCIA_MESES}."""
        veh = vehiculo_completo(
            condCedulaCiudadania="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        doc = self._correr(fake, veh)
        vig = doc.get("estudiosVigencia")
        self.assertIsNotNone(vig)
        delta_meses = ((vig["vence"].year - vig["desde"].year) * 12
                       + vig["vence"].month - vig["desde"].month)
        self.assertEqual(delta_meses, estudios_automaticos.VIGENCIA_MESES)

    # (2026-10-05) Los tests del BARRIDO DE RENOVACIÓN AUTOMÁTICA fueron
    # eliminados junto con la funcionalidad (orden del usuario: la
    # actualización de estudios es MANUAL desde /revision → módulo «Estudios
    # por antigüedad»). El sello de vigencia sigue probándose arriba.

    def test_reintentar_sin_fuentes_fallidas_422(self):
        veh = vehiculo_completo(estudiosSeguridadAuto=[{
            "id": "e1", "tipo": "persona", "cedula": "1020304050",
            "estado": "finalizado", "fuentes": {"SIMIT": False},
            "reporte_id": "rep-1"}])
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake):
            with self.assertRaises(HTTPException) as ctx:
                asyncio.run(estudios_automaticos.reintentar_fuentes_estudio("ABC123", "e1"))
        self.assertEqual(ctx.exception.status_code, 422)

    def _correr_vehiculo(self, fake, vehiculo, forzar=False):
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", _respuesta_ok), \
             patch.object(estudios_automaticos, "consulta_vehiculo", _respuesta_ok), \
             patch.object(estudios_automaticos, "VIGENCIA_MESES", 6):
            asyncio.run(estudios_automaticos.disparar_estudios(
                vehiculo["placa"], forzar=forzar))
        return fake.find_one({"placa": vehiculo["placa"]})

    def test_reutiliza_estudio_de_otra_placa_sin_gastar(self):
        """La misma cédula estudiada hace poco en OTRA placa → se copia el
        resultado (reutilizado_de) sin llamar al proveedor."""
        hace_1_mes = datetime.utcnow() - timedelta(days=30)
        otra = vehiculo_completo(
            placa="OTRA1", condCedulaCiudadania="1020304050",
            estudiosSeguridadAuto=[
                {"id": "x1", "tipo": "persona", "cedula": "1020304050",
                 "estado": "finalizado", "hallazgo": False, "categoria": "",
                 "fuentes": {"SIMIT": False}, "reporte_id": "rep-viejo",
                 "finalizado_en": hace_1_mes},
            ])
        nueva = vehiculo_completo(
            placa="NUEV01", condCedulaCiudadania="1020304050",
            propDocumento="1020304050", tenedDocumento="987654321")
        fake = FakeColeccion([otra, nueva])

        llamadas = {"n": 0}
        async def contar(*a, **kw):
            llamadas["n"] += 1
            return await _respuesta_ok()

        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "VIGENCIA_MESES", 6), \
             patch.object(estudios_automaticos, "consulta_completa", contar), \
             patch.object(estudios_automaticos, "consulta_vehiculo", contar):
            asyncio.run(estudios_automaticos.disparar_estudios("NUEV01"))

        estudios = {e.get("cedula") or e.get("placa"): e
                    for e in fake.find_one({"placa": "NUEV01"})["estudiosSeguridadAuto"]}
        reusado = estudios["1020304050"]
        self.assertEqual(reusado["estado"], "finalizado")
        self.assertEqual(reusado["reutilizado_de"], "OTRA1")
        self.assertEqual(reusado["reporte_id"], "rep-viejo")
        # Solo consultaron los sujetos SIN estudio previo (tenedor + vehículo).
        self.assertEqual(llamadas["n"], 2)

    def test_no_reutiliza_estudios_viejos_ni_con_force(self):
        hace_8_meses = datetime.utcnow() - timedelta(days=245)
        otra = vehiculo_completo(
            placa="OTRA1", condCedulaCiudadania="1020304050",
            estudiosSeguridadAuto=[
                {"id": "x1", "tipo": "persona", "cedula": "1020304050",
                 "estado": "finalizado", "reporte_id": "rep-viejo",
                 "finalizado_en": hace_8_meses},
            ])

        # Viejo (8 meses > vigencia 6): consulta normal.
        nueva = vehiculo_completo(placa="NUEV01", condCedulaCiudadania="1020304050")
        fake = FakeColeccion([otra, nueva])
        doc = self._correr_vehiculo(fake, nueva)
        estudio = [e for e in doc["estudiosSeguridadAuto"]
                   if e["tipo"] == "persona"][0]
        self.assertEqual(estudio["estado"], "finalizado")
        self.assertNotIn("reutilizado_de", estudio)

        # Fresco pero force=True (renovación manual): re-consulta real.
        hace_1_mes = datetime.utcnow() - timedelta(days=30)
        otra["estudiosSeguridadAuto"][0]["finalizado_en"] = hace_1_mes
        fake2 = FakeColeccion([otra, nueva])
        doc2 = self._correr_vehiculo(fake2, nueva, forzar=True)
        estudio2 = [e for e in doc2["estudiosSeguridadAuto"]
                    if e["tipo"] == "persona"][0]
        self.assertNotIn("reutilizado_de", estudio2)

    def test_segunda_corrida_archiva_historial(self):
        """Cada corrida nueva archiva la anterior en `historialEstudios`
        (tope 10) sin perder fechas ni reporte_id."""
        veh = vehiculo_completo(
            condCedulaCiudadasia="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", _respuesta_ok), \
             patch.object(estudios_automaticos, "consulta_vehiculo", _respuesta_ok):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))  # re-disparo manual
        doc = fake.find_one({"placa": "ABC123"})
        self.assertEqual(len(doc["estudiosSeguridadAuto"]), 3)  # corrida vigente nueva
        historial = doc.get("historialEstudios") or []
        self.assertEqual(len(historial), 1)                      # la corrida anterior archivada
        archivados = historial[0]["estudios"]
        self.assertEqual(len(archivados), 3)
        for e in archivados:
            self.assertEqual(e["estado"], "finalizado")
            self.assertEqual(e["reporte_id"], "rep-123")         # abre el PDF histórico

    def test_fallo_de_lanzamiento_reintenta_una_vez(self):
        """El upstream 'realice la consulta nuevamente' (sin jobid) es
        transitorio: se reintenta UNA vez y puede terminar bien."""
        llamadas = {"n": 0}

        async def a_veces_falla(*args, **kwargs):
            llamadas["n"] += 1
            if llamadas["n"] == 1:
                raise HTTPException(status_code=502, detail={
                    "detalle": "TusDatos no entregó jobid",
                    "lanzamiento": {"error": "falla iniciando la consulta"}})
            return await _respuesta_ok()

        veh = vehiculo_completo(
            condCedulaCiudadania="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        import asyncio

        async def _dormir(_segundos):
            return None  # el reintento real espera 5 s; aquí instantáneo

        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", a_veces_falla), \
             patch.object(estudios_automaticos, "consulta_vehiculo", a_veces_falla), \
             patch.object(estudios_automaticos.asyncio, "sleep", _dormir):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
        estudios = fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"]
        # El primer estudio falla el lanzamiento y reintenta (2 llamadas);
        # los otros dos van de una (2+1+1).
        self.assertEqual(llamadas["n"], 4)
        for e in estudios:
            self.assertEqual(e["estado"], "finalizado")

    def test_resultado_no_finalizado_es_error(self):
        async def lenta(*a, **kw):  # _esperar venció el tiempo (_aviso)
            return {"lanzamiento": {"jobid": "j"}, "resultado": {
                "estado": "procesando", "_aviso": "Se agotó el tiempo (300s)"}}

        veh = vehiculo_completo(
            condCedulaCiudadania="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")
        fake = FakeColeccion([veh])
        import asyncio
        with patch.object(estudios_automaticos, "coleccion_vehiculos", fake), \
             patch.object(estudios_automaticos, "_configurado", lambda: True), \
             patch.object(estudios_automaticos, "consulta_completa", lenta), \
             patch.object(estudios_automaticos, "consulta_vehiculo", lenta):
            asyncio.run(estudios_automaticos.disparar_estudios("ABC123"))
        estudios = fake.find_one({"placa": "ABC123"})["estudiosSeguridadAuto"]
        for e in estudios:
            self.assertEqual(e["estado"], "error")
            self.assertIn("agotó el tiempo", e["error"])


class HookActualizarEstadoTests(unittest.TestCase):

    def test_completado_revision_dispara_estudios(self):
        fake = FakeColeccion([vehiculo_completo(
            estadoIntegra="registro_incompleto",
            condCedulaCiudadania="1020304050", propDocumento="1020304050",
            tenedDocumento="987654321")])
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", fake), \
             patch.object(vehiculos, "enviar_notificacion_seguridad") as notif, \
             patch.object(vehiculos, "_disparar_estudios_seguridad") as disparo:
            r = cliente.put("/vehiculos/actualizar-estado", data={
                "placa": "ABC123", "nuevo_estado": "completado_revision",
                "usuario_id": "u1", "nombre_conductor": "Fulano",
            })
        self.assertEqual(r.status_code, 200)
        notif.assert_called_once()
        disparo.assert_called_once_with("ABC123", re_revision=False)

    def test_enviar_a_actualizacion(self):
        """aprobado → en_actualizacion exige observación; finalizar vuelve
        a revisión (y dispara estudios); Seguridad puede cancelar a aprobado."""
        cliente = cliente_de_prueba()

        def llamar(estado_actual, nuevo, extra=None):
            fake = FakeColeccion([vehiculo_completo(estadoIntegra=estado_actual)])
            with patch.object(vehiculos, "coleccion_vehiculos", fake), \
                 patch.object(vehiculos, "enviar_notificacion_seguridad"), \
                 patch.object(vehiculos, "_disparar_estudios_seguridad") as disparo:
                data = {"placa": "ABC123", "nuevo_estado": nuevo, "usuario_id": "u1"}
                data.update(extra or {})
                r = cliente.put("/vehiculos/actualizar-estado", data=data)
            return r, fake, disparo

        # Sin observación → 400 accionable.
        r, _, _ = llamar("aprobado", "en_actualizacion")
        self.assertEqual(r.status_code, 400)

        # Con observación → 200.
        r, fake, _ = llamar("aprobado", "en_actualizacion",
                            {"observaciones": "Actualizar SOAT"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(fake.documents[0]["estadoIntegra"], "en_actualizacion")

        # Conductor finaliza → completado_revision + estudios re-disparados.
        r, fake, disparo = llamar("en_actualizacion", "completado_revision")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(fake.documents[0]["estadoIntegra"], "completado_revision")
        disparo.assert_called_once_with("ABC123", re_revision=False)

        # Seguridad cancela → vuelve a aprobado.
        r, fake, _ = llamar("en_actualizacion", "aprobado",
                            {"motivo": "cancelada"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(fake.documents[0]["estadoIntegra"], "aprobado")

    def test_aprobar_no_dispara_estudios(self):
        # documentosEstudioSeguridad: aprobar exige estudio en PDF (2026-10-06).
        fake = FakeColeccion([vehiculo_completo(
            estadoIntegra="completado_revision",
            documentosEstudioSeguridad=[{"ruta": "Vehiculos/ABC123/estudioSeguridad_test.pdf"}])])
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", fake), \
             patch.object(vehiculos, "_disparar_estudios_seguridad") as disparo:
            r = cliente.put("/vehiculos/actualizar-estado", data={
                "placa": "ABC123", "nuevo_estado": "aprobado", "usuario_id": "u1",
            })
        self.assertEqual(r.status_code, 200)
        disparo.assert_not_called()

    def test_finalizado_por_seguridad_NO_dispara_estudios(self):
        """El alta la finaliza Seguridad (impersonando o directa): los estudios
        los dispara ella con «Volver a consultar» al evaluar la placa — el
        auto-disparo es SOLO del registro hecho por el propio conductor."""
        cliente = cliente_de_prueba()
        for via in ("impersonacion", "seguridad"):
            fake = FakeColeccion([vehiculo_completo(estadoIntegra="registro_incompleto")])
            with patch.object(vehiculos, "coleccion_vehiculos", fake), \
                 patch.object(vehiculos, "enviar_notificacion_seguridad"), \
                 patch.object(vehiculos, "_disparar_estudios_seguridad") as disparo:
                r = cliente.put("/vehiculos/actualizar-estado", data={
                    "placa": "ABC123", "nuevo_estado": "completado_revision",
                    "usuario_id": "u1", "via": via, "editado_por": "SEGURIDAD X",
                })
            self.assertEqual(r.status_code, 200)
            disparo.assert_not_called()


class EndpointEstudiosTests(unittest.TestCase):

    def test_endpoint_devuelve_estudios_del_doc(self):
        fake = FakeColeccion([vehiculo_completo(estudiosSeguridadAuto=[{
            "id": "aaa", "tipo": "persona", "cedula": "1020304050",
            "roles": ["conductor"], "estado": "finalizado",
            "iniciado_en": datetime(2026, 9, 28),
        }])])
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", fake):
            r = cliente.get("/vehiculos/estudios-seguridad/ABC123")
        self.assertEqual(r.status_code, 200)
        cuerpo = r.json()
        self.assertEqual(cuerpo["placa"], "ABC123")
        self.assertEqual(len(cuerpo["estudios"]), 1)
        self.assertEqual(cuerpo["estudios"][0]["estado"], "finalizado")
        self.assertEqual(cuerpo["estudios"][0]["iniciado_en"], "2026-09-28T00:00:00")

    def test_endpoint_devuelve_historico(self):
        fake = FakeColeccion([vehiculo_completo(
            estudiosSeguridadAuto=[{"id": "n1", "tipo": "vehiculo", "placa": "ABC123",
                                    "estado": "en_curso"}],
            historialEstudios=[
                {"fecha": datetime(2026, 9, 20), "estudios": [
                    {"id": "v1", "tipo": "persona", "cedula": "1020304050",
                     "roles": ["conductor"], "estado": "finalizado",
                     "reporte_id": "rep-viejo",
                     "finalizado_en": datetime(2026, 9, 20, 12, 0)},
                ]},
            ])])
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", fake):
            r = cliente.get("/vehiculos/estudios-seguridad/ABC123")
        cuerpo = r.json()
        self.assertEqual(len(cuerpo["estudios"]), 1)      # corrida vigente
        self.assertEqual(len(cuerpo["historico"]), 1)     # corridas anteriores aplanadas
        self.assertEqual(cuerpo["historico"][0]["reporte_id"], "rep-viejo")

    def test_vehiculo_inexistente_404(self):
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", FakeColeccion()):
            r = cliente.get("/vehiculos/estudios-seguridad/NADA")
        self.assertEqual(r.status_code, 404)

    def test_exportar_excel(self):
        """GET /exportar-excel: todas las placas, columnas clave presentes."""
        fake = FakeColeccion([
            vehiculo_completo(placa="AAA111", estadoIntegra="aprobado",
                               condCedulaCiudadaria="1020304050",
                               condNombres="JUAN", estudiosVigencia={
                                   "desde": datetime(2026, 9, 28),
                                   "vence": datetime(2027, 9, 28)}),
            vehiculo_completo(placa="ZZZ999", estadoIntegra="registro_incompleto"),
        ])
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", fake):
            r = cliente.get("/vehiculos/exportar-excel")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        self.assertIn("attachment", r.headers.get("content-disposition", ""))
        # Verificar contenido del xlsx (openpyxl en memoria).
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        cabeceras = [c.value for c in ws[1]]
        self.assertIn("Placa", cabeceras)
        self.assertIn("Cond. Cédula", cabeceras)
        self.assertIn("Estudios vigencia vence", cabeceras)
        self.assertEqual(ws.max_row, 3)  # cabecera + 2 vehículos
        placas = [ws.cell(row=i, column=1).value for i in (2, 3)]
        self.assertEqual(placas, ["AAA111", "ZZZ999"])  # ordenadas

    def test_disparar_manual(self):
        """POST .../disparar: botón de /revision — re-dispara SIEMPRE con
        force (re-consulta real, no la caché del proveedor) y manual=True
        (2026-10-07: el disparo MANUAL ignora el switch de disparo
        automático — antes respondía 200 en silencio con el switch apagado)."""
        fake = FakeColeccion([vehiculo_completo()])
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", fake), \
             patch.object(vehiculos, "_disparar_estudios_seguridad") as disparo:
            r = cliente.post("/vehiculos/estudios-seguridad/ABC123/disparar")
        self.assertEqual(r.status_code, 200)
        disparo.assert_called_once_with("ABC123", re_revision=False, forzar=True,
                                        manual=True)

    def test_disparar_vehiculo_inexistente_404(self):
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", FakeColeccion()), \
             patch.object(vehiculos, "_disparar_estudios_seguridad") as disparo:
            r = cliente.post("/vehiculos/estudios-seguridad/NADA/disparar")
        self.assertEqual(r.status_code, 404)
        disparo.assert_not_called()


class BlindajesTests(unittest.TestCase):

    def test_campos_blindados(self):
        self.assertIn("estudiosSeguridadAuto", vehiculos.CLAVES_PROTEGIDAS)
        self.assertIn("historialEstudios", vehiculos.CLAVES_PROTEGIDAS)
        self.assertIn("estudiosSeguridadAuto", vehiculos.CAMPOS_VOLATILES_FIRMA)
        self.assertIn("historialEstudios", vehiculos.CAMPOS_VOLATILES_FIRMA)

    def test_listados_ocultan_estudios_al_conductor(self):
        fake = FakeColeccion([vehiculo_completo(
            estadoIntegra="completado_revision",
            estudiosSeguridadAuto=[{"id": "x", "estado": "finalizado"}],
            historialEstudios=[{"fecha": datetime(2026, 9, 20), "estudios": []}])])
        cliente = cliente_de_prueba()
        with patch.object(vehiculos, "coleccion_vehiculos", fake), \
             patch.object(vehiculos, "_firmar_documentos", lambda v: v):
            r = cliente.get("/vehiculos/obtener-vehiculos-incompletos")
        self.assertEqual(r.status_code, 200)
        for v in r.json()["vehicles"]:
            self.assertNotIn("estudiosSeguridadAuto", v)
            self.assertNotIn("historialEstudios", v)


if __name__ == "__main__":
    unittest.main()
